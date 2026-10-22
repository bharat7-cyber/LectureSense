from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

import openpyxl


def parse_num(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    # Handles "1.07s", "00:01:02.5", and plain numbers.
    if text.lower().endswith("s"):
        text = text[:-1].strip()
    if ":" in text:
        parts = text.split(":")
        try:
            nums = [float(x) for x in parts]
            if len(nums) == 3:
                return nums[0] * 3600 + nums[1] * 60 + nums[2]
            if len(nums) == 2:
                return nums[0] * 60 + nums[1]
        except ValueError:
            pass
    try:
        return float(text)
    except ValueError:
        return None


def load_timeline(path: Path):
    data = json.loads(path.read_text(encoding="utf-8"))
    timeline = data.get("timeline", data if isinstance(data, list) else [])
    if not isinstance(timeline, list):
        raise ValueError("Timeline JSON does not contain a 'timeline' list.")

    entries = []
    for item in timeline:
        sid = item.get("slide_id", item.get("slide"))
        start = parse_num(item.get("start_time", item.get("start")))
        end = parse_num(item.get("end_time", item.get("end")))
        if sid is None or start is None or end is None:
            continue
        try:
            sid = int(sid)
        except (TypeError, ValueError):
            continue
        entries.append((start, end, sid))

    if not entries:
        raise ValueError(f"No usable timeline entries in {path}")

    entries.sort(key=lambda x: (x[0], x[1], x[2]))
    return entries


def load_ground_truth(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Ground-truth workbook is empty: {path}")

    headers = [str(x).strip().lower() if x is not None else "" for x in rows[0]]

    def find_col(*names):
        for name in names:
            name = name.lower()
            if name in headers:
                return headers.index(name)
        return None

    time_col = find_col("time", "timestamp", "start", "key")
    slide_col = find_col("slidenumber", "slide_number", "slide", "slideid")

    if time_col is None or slide_col is None:
        raise ValueError(
            f"Could not find Time/Slidenumber columns in {path}. "
            f"Headers found: {headers}"
        )

    gt = []
    for row in rows[1:]:
        if time_col >= len(row) or slide_col >= len(row):
            continue
        t = parse_num(row[time_col])
        s = parse_num(row[slide_col])
        if t is None or s is None:
            continue
        s = int(s)
        # MaViLS uses -1 for points where a unique slide cannot be identified.
        if s < 0:
            continue
        gt.append((t, s))

    wb.close()
    return gt


def predict_slide(entries, t):
    # A timestamp belongs to the interval [start, end]. In the rare case
    # of overlapping intervals, use the interval with the latest start.
    matches = [e for e in entries if e[0] <= t <= e[1]]
    if matches:
        return max(matches, key=lambda e: (e[0], e[1]))[2]

    # For small gaps between generated timeline entries, assign the
    # nearest interval boundary. This is only relevant when the timestamp
    # falls in a gap not covered by any predicted segment.
    before = [e for e in entries if e[1] < t]
    after = [e for e in entries if e[0] > t]

    candidates = []
    if before:
        e = max(before, key=lambda x: x[1])
        candidates.append((abs(t - e[1]), e[2]))
    if after:
        e = min(after, key=lambda x: x[0])
        candidates.append((abs(e[0] - t), e[2]))

    if not candidates:
        return None
    return min(candidates, key=lambda x: x[0])[1]


def evaluate(timeline_path: Path, gt_path: Path, verbose=True):
    entries = load_timeline(timeline_path)
    gt = load_ground_truth(gt_path)

    correct = 0
    errors = []

    for t, gt_slide in gt:
        pred = predict_slide(entries, t)
        ok = pred == gt_slide
        if ok:
            correct += 1
        else:
            errors.append((t, gt_slide, pred))

        if verbose:
            status = "CORRECT" if ok else "WRONG"
            pred_text = "NONE" if pred is None else str(pred)
            print(f"{t:8.2f}s | GT Slide {gt_slide:2d} | Pred Slide {pred_text:>2} | {status}")

    total = len(gt)
    accuracy = correct / total if total else 0.0

    print("\n" + "=" * 70)
    print("SUMMARY")
    print("=" * 70)
    print(f"Correct : {correct}/{total}")
    print(f"Accuracy: {accuracy:.4f} ({accuracy:.2%})")
    print(f"Timeline entries: {len(entries)}")

    if errors:
        print("\nErrors:")
        for t, gt_slide, pred in errors:
            pred_text = "NONE" if pred is None else str(pred)
            print(f"  {t:.2f}s | GT {gt_slide} | Pred {pred_text}")

    return correct, total, accuracy, errors


def main():
    parser = argparse.ArgumentParser(
        description="Correct MaViLS timeline evaluator. Unlike the old evaluator, "
                    "this script actually uses the supplied timeline and GT paths."
    )
    parser.add_argument("--timeline", required=True)
    parser.add_argument("--ground-truth", required=True)
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args()

    timeline = Path(args.timeline)
    gt = Path(args.ground_truth)

    if not timeline.exists():
        raise FileNotFoundError(f"Timeline not found: {timeline}")
    if not gt.exists():
        raise FileNotFoundError(f"Ground truth not found: {gt}")

    print("=" * 70)
    print("GROUND TRUTH VS TIMELINE")
    print("=" * 70)
    print(f"Timeline:     {timeline}")
    print(f"Ground truth: {gt}")
    print()

    evaluate(timeline, gt, verbose=not args.quiet)


if __name__ == "__main__":
    main()
